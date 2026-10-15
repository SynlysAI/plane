# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only

"""Contract coverage for plain workspace member invitation imports."""

import io

import pytest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from plane.db.models import (
    MentorBinding,
    OrgUnitMember,
    ResearchUserProfile,
    UserImportBatch,
    WorkspaceMember,
    WorkspaceMemberInvite,
)
from plane.research.utils.capabilities import research_level
from plane.tests.research_fixtures import make_user, make_workspace


pytestmark = pytest.mark.contract


def workbook_bytes(rows):
    """Build an in-memory member invitation workbook."""
    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


def client_for(user):
    """Return a client authenticated as the supplied user."""
    client = APIClient()
    client.force_authenticate(user=user)
    return client


@pytest.fixture
def env(db):
    """Create an administrator and a plain workspace."""
    admin = make_user(email="admin@example.com", first_name="Admin")
    workspace = make_workspace(admin, name="Plain Workspace", slug="plain-ws")
    return {"admin": admin, "workspace": workspace}


def template_url(workspace):
    """Return the plain invitation template URL."""
    return f"/api/workspaces/{workspace.slug}/invitations/import/template/"


def import_url(workspace):
    """Return the plain invitation import URL."""
    return f"/api/workspaces/{workspace.slug}/invitations/import/"


def test_template_contains_only_plain_member_columns(env):
    response = client_for(env["admin"]).get(template_url(env["workspace"]))

    assert response.status_code == 200
    worksheet = load_workbook(io.BytesIO(response.content), read_only=True).active
    assert [cell.value for cell in next(worksheet.iter_rows(max_row=1))] == ["名称", "邮箱", "角色"]


def test_valid_workbook_creates_pending_invitations_without_research_relations(env):
    payload = workbook_bytes(
        [
            ["邮箱", "名称", "角色"],
            ["member@example.com", "普通成员", "成员"],
        ]
    )

    response = client_for(env["admin"]).post(
        import_url(env["workspace"]),
        {
            "file": io.BytesIO(payload),
            "research_profile_enabled": "true",
            "research_profile_category": "STUDENT",
        },
        format="multipart",
    )

    assert response.status_code == 200
    assert response.data["created"] == 1
    invitation = WorkspaceMemberInvite.objects.get(workspace=env["workspace"], email="member@example.com")
    assert invitation.display_name == "普通成员"
    assert invitation.role == 15
    assert invitation.research_profile_enabled is True
    assert invitation.research_profile_category == "STUDENT"
    assert not UserImportBatch.objects.exists()
    assert not OrgUnitMember.objects.exists()
    assert not MentorBinding.objects.exists()


def test_invalid_workbook_is_rejected_as_a_whole(env):
    payload = workbook_bytes(
        [
            ["名称", "邮箱", "角色"],
            ["张三", "duplicate@example.com", "成员"],
            ["李四", "duplicate@example.com", "访客"],
        ]
    )

    response = client_for(env["admin"]).post(
        import_url(env["workspace"]),
        {"file": io.BytesIO(payload)},
        format="multipart",
    )

    assert response.status_code == 400
    assert response.data["rows"][0]["row_number"] == 3
    assert "邮箱重复" in response.data["rows"][0]["message"]
    assert not WorkspaceMemberInvite.objects.exists()


def test_pending_invitation_duplicate_is_rejected_without_a_second_invite(env):
    WorkspaceMemberInvite.objects.create(
        workspace=env["workspace"],
        email="pending@example.com",
        token="existing-token",
        role=5,
    )
    payload = workbook_bytes(
        [
            ["名称", "邮箱", "角色"],
            ["待邀请", "pending@example.com", "成员"],
        ]
    )

    response = client_for(env["admin"]).post(
        import_url(env["workspace"]),
        {"file": io.BytesIO(payload)},
        format="multipart",
    )

    assert response.status_code == 400
    assert "already has a pending invitation" in response.data["rows"][0]["message"]
    assert WorkspaceMemberInvite.objects.filter(email="pending@example.com").count() == 1


def test_accepting_basic_profile_invitation_creates_profile_without_org_relation(env):
    payload = workbook_bytes(
        [
            ["名称", "邮箱", "角色"],
            ["学生", "student@example.com", "成员"],
        ]
    )
    admin_client = client_for(env["admin"])
    assert (
        admin_client.post(
            import_url(env["workspace"]),
            {
                "file": io.BytesIO(payload),
                "research_profile_enabled": True,
                "research_profile_category": "STUDENT",
            },
            format="multipart",
        ).status_code
        == 200
    )
    invitation = WorkspaceMemberInvite.objects.get(email="student@example.com")
    invited = make_user(email="student@example.com", first_name="Student")

    with patch("plane.db.mixins.soft_delete_related_objects.delay"):
        response = client_for(invited).post(
            f"/api/workspaces/{env['workspace'].slug}/invitations/{invitation.id}/join/",
            {"token": invitation.token, "accepted": True},
            format="json",
        )

    assert response.status_code == 200
    assert WorkspaceMember.objects.filter(workspace=env["workspace"], member=invited, role=15).exists()
    assert ResearchUserProfile.objects.filter(user=invited, category="STUDENT").exists()
    assert not OrgUnitMember.objects.filter(user=invited).exists()
    assert research_level(invited, env["workspace"]) == "NONE"


def test_basic_profile_invite_does_not_overwrite_an_existing_profile(env):
    invited = make_user(email="advisor@example.com", first_name="Advisor")
    ResearchUserProfile.objects.create(user=invited, category="ADVISOR")
    invitation = WorkspaceMemberInvite.objects.create(
        workspace=env["workspace"],
        email=invited.email,
        token="invite-token",
        role=15,
        research_profile_enabled=True,
        research_profile_category="STUDENT",
    )

    with patch("plane.db.mixins.soft_delete_related_objects.delay"):
        response = client_for(invited).post(
            f"/api/workspaces/{env['workspace'].slug}/invitations/{invitation.id}/join/",
            {"token": invitation.token, "accepted": True},
            format="json",
        )

    assert response.status_code == 200
    invited.research_profile.refresh_from_db()
    assert invited.research_profile.category == "ADVISOR"


def test_bulk_acceptance_applies_basic_profile_options(env):
    invited = make_user(email="bulk@example.com", first_name="Bulk")
    invitation = WorkspaceMemberInvite.objects.create(
        workspace=env["workspace"],
        email=invited.email,
        token="bulk-token",
        role=5,
        research_profile_enabled=True,
        research_profile_category="OTHER",
    )

    response = client_for(invited).post(
        "/api/users/me/workspaces/invitations/",
        {"invitations": [str(invitation.id)]},
        format="json",
    )

    assert response.status_code == 204
    assert WorkspaceMember.objects.filter(workspace=env["workspace"], member=invited, role=5).exists()
    assert ResearchUserProfile.objects.filter(user=invited, category="OTHER").exists()


def test_legacy_invitation_payload_remains_compatible(env):
    response = client_for(env["admin"]).post(
        f"/api/workspaces/{env['workspace'].slug}/invitations/",
        {"emails": [{"email": "legacy@example.com", "role": 5}]},
        format="json",
    )

    assert response.status_code == 200
    invitation = WorkspaceMemberInvite.objects.get(email="legacy@example.com")
    assert invitation.display_name == ""
    assert invitation.research_profile_enabled is False
    assert invitation.research_profile_category == ""
