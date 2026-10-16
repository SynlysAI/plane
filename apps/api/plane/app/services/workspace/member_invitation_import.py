# Copyright (c) 2023-present Plane Software, Inc. and contributors
# SPDX-License-Identifier: AGPL-3.0-only

"""Parsing helpers for plain workspace member invitation uploads."""

import csv
import io
from dataclasses import dataclass

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

from openpyxl import Workbook, load_workbook


MEMBER_IMPORT_HEADERS = {
    "name": "名称",
    "email": "邮箱",
    "role": "角色",
}

ROLE_ALIASES = {
    "20": 20,
    "管理员": 20,
    "admin": 20,
    "15": 15,
    "成员": 15,
    "member": 15,
    "5": 5,
    "访客": 5,
    "guest": 5,
}


@dataclass(frozen=True)
class MemberInvitationRow:
    """One validated candidate for a plain workspace invitation."""

    row_number: int
    display_name: str
    email: str
    role: int
    research_profile_enabled: bool = False
    research_profile_category: str = ""


class MemberInvitationImportError(Exception):
    """Raised when an invitation workbook cannot be used as a whole batch."""

    def __init__(self, message, rows=None):
        self.message = message
        self.rows = rows or []
        super().__init__(message)


def normalise_member_role(value):
    """Convert a UI, JSON, CSV, or Excel role value to a workspace role number.

    Args:
        value: The raw role value from an upload or request payload.

    Returns:
        The role integer, or ``None`` when the value is not a supported alias.
    """
    raw = str(value if value is not None else "").strip()
    if raw.endswith(".0"):
        raw = raw[:-2]
    return ROLE_ALIASES.get(raw.lower())


def normalise_boolean(value):
    """Convert JSON booleans and form-data strings to a boolean value.

    Args:
        value: The raw request value.

    Returns:
        The normalized boolean.
    """
    if isinstance(value, str):
        return value.strip().lower() in {"1", "true", "yes", "on"}
    return bool(value)


def build_member_invitation_template():
    """Build a three-column XLSX template for plain workspace invitations.

    Returns:
        The workbook bytes ready to be returned as an HTTP response.
    """
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "成员邀请"
    worksheet.append(["名称", "邮箱", "角色"])
    worksheet.column_dimensions["A"].width = 24
    worksheet.column_dimensions["B"].width = 32
    worksheet.column_dimensions["C"].width = 18
    worksheet.freeze_panes = "A2"
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.read()


def _normalise_header(value):
    """Normalize a worksheet header for tolerant Chinese header matching."""
    return "".join(str(value or "").split()).lower()


def _decode_csv(payload):
    """Decode a CSV upload while accepting common Chinese encodings."""
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise MemberInvitationImportError("CSV 编码必须是 UTF-8 或 GB18030。")


def _sheet_rows(payload, filename):
    """Read the first worksheet of an XLSX or CSV upload into rows."""
    if not payload:
        raise MemberInvitationImportError("请选择要导入的 Excel 或 CSV 文件。")
    lowered = str(filename or "").lower()
    try:
        if lowered.endswith(".xlsx") or payload[:2] == b"PK":
            workbook = load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
            worksheet = workbook[workbook.sheetnames[0]]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            workbook.close()
            return rows
        rows = [list(row) for row in csv.reader(io.StringIO(_decode_csv(payload)))]
        return rows
    except MemberInvitationImportError:
        raise
    except Exception as exc:  # noqa: BLE001 - converted to a user-facing 400
        raise MemberInvitationImportError(f"文件无法解析：{exc}") from exc


def _map_headers(header_row):
    """Map canonical field names to header positions."""
    normalized = [_normalise_header(cell) for cell in header_row]
    non_empty = [value for value in normalized if value]
    if len(non_empty) != len(set(non_empty)):
        raise MemberInvitationImportError("表头存在重复列，请检查列名。")
    mapping = {}
    for field_name, header in MEMBER_IMPORT_HEADERS.items():
        normalized_header = _normalise_header(header)
        if normalized_header not in normalized:
            raise MemberInvitationImportError("表头必须包含：名称、邮箱、角色。")
        mapping[field_name] = normalized.index(normalized_header)
    return mapping


def parse_member_invitation_workbook(
    payload, filename, *, research_profile_enabled=False, research_profile_category=""
):
    """Parse and validate a plain workspace member invitation workbook.

    Args:
        payload: The uploaded file bytes.
        filename: The uploaded filename, used to select XLSX or CSV parsing.
        research_profile_enabled: Whether all imported members get a basic profile.
        research_profile_category: The basic profile category applied to all rows.

    Returns:
        A list of validated :class:`MemberInvitationRow` objects.

    Raises:
        MemberInvitationImportError: When any row is invalid or duplicated.
    """
    rows = _sheet_rows(payload, filename)
    if not rows:
        raise MemberInvitationImportError("文件为空。")
    mapping = _map_headers(rows[0])
    errors = []
    seen_emails = {}
    parsed = []

    def cell(row, field_name):
        position = mapping[field_name]
        return str(row[position]).strip() if position < len(row) and row[position] is not None else ""

    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(value or "").strip() for value in row):
            continue
        display_name = cell(row, "name")
        email = cell(row, "email").lower()
        role = normalise_member_role(cell(row, "role"))
        message = ""
        if not display_name:
            message = "名称不能为空。"
        elif len(display_name) > 255:
            message = "名称不能超过 255 个字符。"
        elif not email:
            message = "邮箱不能为空。"
        elif len(email) > 255:
            message = "邮箱不能超过 255 个字符。"
        else:
            try:
                validate_email(email)
            except ValidationError:
                message = f"{email} 不是有效邮箱。"
        if not message and role is None:
            message = "角色必须是管理员、成员或访客。"
        if not message and email in seen_emails:
            message = f"邮箱重复于第 {seen_emails[email]} 行。"
        if message:
            errors.append({"row_number": row_number, "name": display_name, "email": email, "message": message})
            continue
        seen_emails[email] = row_number
        parsed.append(
            MemberInvitationRow(
                row_number=row_number,
                display_name=display_name,
                email=email,
                role=role,
                research_profile_enabled=research_profile_enabled,
                research_profile_category=research_profile_category,
            )
        )

    if not parsed and not errors:
        raise MemberInvitationImportError("文件只有表头，没有数据行。")
    if errors:
        raise MemberInvitationImportError("成员导入表存在无效行。", rows=errors)
    return parsed
